from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = Document()

# Margins
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

# Default font
styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(11)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Reinforcement Learning – Exam-Ready Answers")
run.bold = True
run.font.size = Pt(16)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("Important Questions and Answers")
r.italic = True
r.font.size = Pt(10)

def heading(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13)
    return p

def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p

def numbered(text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p

# Q1
heading("1. Three Primary Components of a Reinforcement Learning System")
p = doc.add_paragraph("The three primary components are:")
numbered("Agent – The learner or decision-maker that interacts with the environment.")
numbered("Environment – The external system with which the agent interacts.")
numbered("Reward – A numerical feedback signal that tells the agent how good or bad its action was.")

p = doc.add_paragraph()
r = p.add_run("Interaction in a single time step")
r.bold = True

steps = [
    "The agent observes the current state Sₜ.",
    "Based on its policy, the agent selects an action Aₜ.",
    "The environment receives the action and changes to a new state Sₜ₊₁.",
    "The environment gives the agent a reward Rₜ₊₁.",
    "The agent uses this experience to improve its future decisions."
]
for s in steps:
    numbered(s)

p = doc.add_paragraph()
r = p.add_run("Interaction: ")
r.bold = True
p.add_run("Sₜ → Aₜ → Rₜ₊₁, Sₜ₊₁")

p = doc.add_paragraph()
p.add_run("Example: ").bold = True
p.add_run("In a game, the agent is the player, the environment is the game, and the reward may be +10 for winning and −10 for losing.")

# Q2
heading("2. Policy and Value Function")
p = doc.add_paragraph()
p.add_run("Policy: ").bold = True
p.add_run("A policy defines how an agent chooses an action when it is in a particular state.")

p = doc.add_paragraph()
p.add_run("It is usually represented as: ").bold = False
p.add_run("π(a|s)")

p = doc.add_paragraph()
p.add_run("Value Function: ").bold = True
p.add_run("A value function estimates how good it is for an agent to be in a particular state, considering the expected future rewards.")

p = doc.add_paragraph()
p.add_run("State-value function: ")
p.add_run("V^π(s) = E_π[Gₜ | Sₜ = s]")

p = doc.add_paragraph()
p.add_run("Difference:")
p.runs[0].bold = True

table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text = "Policy"
hdr[1].text = "Value Function"
rows = [
    ("Determines what action to take", "Determines how good a state is"),
    ("Guides decision-making", "Evaluates future rewards"),
    ("Directly selects actions", "Helps compare states/actions"),
    ('Example: "Move left"', 'Example: "This state is worth 8 points"')
]
for a, b in rows:
    cells = table.add_row().cells
    cells[0].text = a
    cells[1].text = b

p = doc.add_paragraph()
p.add_run("Why both are important: ").bold = True
p.add_run("The policy tells the agent what to do, while the value function tells it how good the consequences are. Together, they help the agent improve its decisions and move toward an optimal policy that maximizes cumulative reward.")

# Q3
heading("3. Exploration vs Exploitation")
p = doc.add_paragraph()
p.add_run("Exploration: ").bold = True
p.add_run("Trying new or unfamiliar actions to discover whether they can produce better rewards.")

p = doc.add_paragraph()
p.add_run("Example: ").bold = True
p.add_run("A robot tries a new path even though it already knows one path that works.")

p = doc.add_paragraph()
p.add_run("Exploitation: ").bold = True
p.add_run("Choosing the action that the agent currently believes will give the highest reward.")

p = doc.add_paragraph()
p.add_run("Example: ").bold = True
p.add_run("The robot repeatedly uses the path it currently considers the shortest.")

p = doc.add_paragraph()
p.add_run("Difference:")
p.runs[0].bold = True

table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text = "Exploration"
hdr[1].text = "Exploitation"
rows = [
    ("Tries new actions", "Uses known best actions"),
    ("Discovers new possibilities", "Uses existing knowledge"),
    ("May give lower immediate reward", "Usually gives higher immediate reward")
]
for a, b in rows:
    cells = table.add_row().cells
    cells[0].text = a
    cells[1].text = b

p = doc.add_paragraph()
p.add_run("Why balance is important: ").bold = True
p.add_run("If an agent only explores, it may waste time trying poor actions. If it only exploits, it may never discover a better strategy. Therefore, a proper balance is necessary for successful learning and achieving high long-term rewards.")

p = doc.add_paragraph()
p.add_run("Common method – Epsilon-Greedy: ").bold = True
p.add_run("With probability ε, the agent chooses a random action (exploration). With probability 1 − ε, it chooses the best-known action (exploitation).")

p = doc.add_paragraph()
p.add_run("Example: ").bold = True
p.add_run("If ε = 0.1, the agent explores 10% of the time and exploits 90% of the time. Usually, ε is gradually decreased during training.")

# Q4
heading("4. Markov Decision Process (MDP)")
p = doc.add_paragraph()
p.add_run("A reinforcement learning problem is commonly represented using a Markov Decision Process (MDP). An MDP consists of four main elements:")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("MDP = (S, A, P, R)")
r.bold = True

items = [
    ("1. State (S)", "The state represents the current situation of the environment. Example: In a chess game, the current arrangement of all pieces represents the state."),
    ("2. Actions (A)", "The actions are the possible choices available to the agent. Example: In chess, the possible legal moves are the actions."),
    ("3. Transition Probability (P)", "The transition function describes the probability of moving from one state to another after taking an action. P(s′|s,a) represents the probability of reaching state s′ when action a is taken in state s."),
    ("4. Reward (R)", "The reward function specifies the immediate numerical feedback received after taking an action. It may be +10 for winning, −10 for losing, or 0 for a normal move.")
]
for title_text, body in items:
    p = doc.add_paragraph()
    p.add_run(title_text + ": ").bold = True
    p.add_run(body.split(": ", 1)[-1] if ": " in body else body)

p = doc.add_paragraph()
p.add_run("Markov Property: ").bold = True
p.add_run("The Markov property states that the future depends only on the current state and action, and not on the complete history of previous states.")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("P(Sₜ₊₁ | Sₜ, Aₜ, Sₜ₋₁, Aₜ₋₁, …) = P(Sₜ₊₁ | Sₜ, Aₜ)").bold = True

p = doc.add_paragraph()
p.add_run("In simple words: ").bold = True
p.add_run("The current state contains all the necessary information needed to predict the future.")

p = doc.add_paragraph()
p.add_run("Significance: ").bold = True
p.add_run("The Markov property makes RL problems easier to model because the agent does not need to remember the entire history of interactions. It only needs the current state to make decisions.")

# Quick revision
heading("Quick Revision")
for text in [
    "RL interaction: State → Action → Reward + New State",
    "Policy: What should I do?",
    "Value function: How good is this state/action?",
    "Exploration: Try something new.",
    "Exploitation: Use what I already know.",
    "MDP: (S, A, P, R)",
    "Markov property: The future depends on the present, not the entire past."
]:
    bullet(text)

# Footer
for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Reinforcement Learning – Exam Notes").font.size = Pt(9)

path = "/mnt/data/Reinforcement_Learning_Exam_Answers.docx"
doc.save(path)
print(path)
